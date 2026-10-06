"""Check a full-format recalculation against the frozen source and record it.
Usage: python full_source_check.py SOURCE RECALCULATED AUDIT_DIRECTORY
"""
from pathlib import Path
import hashlib,json,math,re,shutil,sys
import openpyxl

def main(source,calculated,audit):
    source,calculated,audit=map(Path,[source,calculated,audit])
    w=openpyxl.load_workbook(source);r=openpyxl.load_workbook(calculated);d=openpyxl.load_workbook(calculated,data_only=True)
    checks=[]
    def expect(name,test,detail=None):checks.append({'name':name,'pass':bool(test),'detail':detail})
    def normal(v):return re.sub(r'\b(TRUE|FALSE)\(\)',r'\1',v) if isinstance(v,str) and v.startswith('=') else v
    diff=[(s.title,c.coordinate) for s in w for row in s for c in row if normal(c.value)!=normal(r[s.title][c.coordinate].value)]
    expect('same 19 modules',w.sheetnames==r.sheetnames and len(w.sheetnames)==19)
    expect('every input and formula preserved',not diff,diff)
    errors=[(s.title,c.coordinate,c.value) for s in d for row in s for c in row if c.data_type=='e']
    expect('full workbook has no formula errors',not errors,errors)
    expect('full workbook system PASS',d['SYSTEM CHECK']['B19'].value=='PASS')
    expect('full workbook initial decision GO',d['DECISION CENTER']['B16'].value=='GO')
    expect('version is v2.2',str(d['SETUP']['B24'].value)=='2.2')
    expect('numeric values are finite',all(math.isfinite(c.value) for s in d for row in s for c in row if isinstance(c.value,(int,float))))
    baseline=openpyxl.load_workbook(audit/'recalculated/baseline.xlsx',data_only=True)
    mismatches=[]
    for s in w:
        for row in s:
            for c in row:
                if c.data_type!='f':continue
                a=d[s.title][c.coordinate].value;b=baseline[s.title][c.coordinate].value
                equal=math.isclose(a,b,rel_tol=1e-12,abs_tol=1e-12) if isinstance(a,(int,float)) and isinstance(b,(int,float)) else a==b
                if not equal:mismatches.append((s.title,c.coordinate,a,b))
    expect('full-format and lightweight baselines agree at every formula',not mismatches,mismatches)
    receipt={'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'recalculated_sha256':hashlib.sha256(calculated.read_bytes()).hexdigest(),'engine':'LibreOffice headless; complete original presentation retained on input','checks':checks,'status':'PASS' if all(x['pass'] for x in checks) else 'FAIL'}
    (audit/'full_source_receipt.json').write_text(json.dumps(receipt,indent=2))
    if receipt['status']=='PASS':shutil.copy2(calculated,audit/'full_source_recalculated.xlsx')
    print(json.dumps(receipt));return receipt['status']!='PASS'

if __name__=='__main__':sys.exit(main(*sys.argv[1:]))
