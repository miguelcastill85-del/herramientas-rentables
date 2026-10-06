"""Package the audited correction and refreshed guide; retain original XLSX layout.

Usage: python finalize_package.py V2_1_PACKAGE_DIR V2_2_XLSX AUDIT_DIR OUTPUT_DIR
"""
from pathlib import Path
from xml.etree import ElementTree as ET
import hashlib, json, re, sys, zipfile
import openpyxl
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import Paragraph, Spacer
from reportlab.pdfgen import canvas
from reportlab.lib.utils import simpleSplit
from xml.sax.saxutils import escape

def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()

def cache_into_original(source,calculated,target):
    """Copy recalculated values, not engine-exported formatting or formula rewrites."""
    formula=openpyxl.load_workbook(source)
    data=openpyxl.load_workbook(calculated,data_only=True)
    calc_formulas=openpyxl.load_workbook(calculated)
    for sn in formula.sheetnames:
        for row in formula[sn]:
            for c in row:
                if c.data_type=='f':assert calc_formulas[sn][c.coordinate].data_type=='f'
    ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    ET.register_namespace('',ns)
    with zipfile.ZipFile(source) as src,zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as dest:
        for info in src.infolist():
            raw=src.read(info.filename)
            m=re.fullmatch(r'xl/worksheets/sheet(\d+)\.xml',info.filename)
            if m:
                sn=formula.sheetnames[int(m[1])-1]; root=ET.fromstring(raw)
                for cell in root.findall('.//{'+ns+'}c'):
                    if cell.find('{'+ns+'}f') is None:continue
                    value=data[sn][cell.attrib['r']].value
                    old=cell.find('{'+ns+'}v')
                    if old is not None:cell.remove(old)
                    v=ET.SubElement(cell,'{'+ns+'}v')
                    if value is None:cell.set('t','str');v.text=''
                    elif isinstance(value,bool):cell.set('t','b');v.text='1' if value else '0'
                    elif isinstance(value,(int,float)):
                        cell.attrib.pop('t',None);v.text=str(value)
                    else:cell.set('t','str');v.text=str(value)
                raw=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            dest.writestr(info,raw)
    actual=openpyxl.load_workbook(target)
    for sn in formula.sheetnames:
        for row in formula[sn]:
            for c in row:assert actual[sn][c.coordinate].value==c.value

def build_guide(old,target):
    reader=PdfReader(old)
    assert len(reader.pages)==9
    extras={0:'v2.2 independent audit correction: complete project formulas, fail-closed input checks, no circular model-health calculation, valid historical samples, and cost-only scope analysis.',
      1:'Defaults are illustrative, not recommended tax reserves or payment-channel fees. Replace them with your actual values. The storefront price in CLP is separate from the currency used for your own projects.',
      4:'For a costs-only change, the effective rate also falls even when no extra hours are added. Additional fees use the selected complexity and fee assumptions; fixed fees must reflect the actual number of charges.',
      7:'Export only CLIENT SUMMARY as PDF. Do not send the internal workbook to a client. The .xlsx contains private assumptions and project records.',
      8:'Enter explicit zero for no external costs, fees or hidden hours. Missing values are incomplete evidence. Historical confidence counts usable hours records, not every row marked Completed. Formula cells have password-free protection against accidental edits; use Review > Unprotect Sheet for deliberate customization.'}
    c=canvas.Canvas(str(target),pagesize=(612,792));c.setTitle('Freelancer Pricing Intelligence System v2.2 - Quick Start')
    style=ParagraphStyle('body',fontName='Helvetica',fontSize=11,leading=16,textColor=colors.HexColor('#213044'))
    for i,page in enumerate(reader.pages):
        text=page.extract_text().replace('v2.1','v2.2')
        text=re.sub(r'Herramientas Rentables · DP-FREELANCE-001 · v2.2\s*Page \d+','',text).strip()
        if i==0:
            text=text.split('v2.2 quality patch')[0].strip()
        lines=text.splitlines();title=lines[0]
        c.setFillColor(colors.HexColor('#0b1b31'));c.rect(0,700,612,92,fill=1,stroke=0)
        c.setFillColor(colors.white);c.setFont('Helvetica-Bold',18)
        for j,line in enumerate(simpleSplit(title,'Helvetica-Bold',18,510)):c.drawString(46,757-23*j,line)
        c.setFont('Helvetica',10);c.drawString(46,719,'HERRAMIENTAS RENTABLES  /  DP-FREELANCE-001  /  v2.2')
        body=' '.join(lines[1:])
        p=Paragraph(escape(body),style);_,h=p.wrap(520,610);p.drawOn(c,46,675-h)
        if i in extras:
            extra=Paragraph(escape(extras[i]),style);_,eh=extra.wrap(490,300)
            y=675-h-34-eh
            assert y>65,'Guide content would overflow'
            c.setFillColor(colors.HexColor('#e8f4f7'));c.roundRect(40,y-16,532,eh+32,8,fill=1,stroke=0)
            extra.drawOn(c,58,y)
        c.setFillColor(colors.HexColor('#526174'));c.setFont('Helvetica',9)
        c.drawString(46,34,'Business planning tool. No guaranteed income. English content.')
        c.drawRightString(566,34,f'{i+1} / 9');c.showPage()
    c.save()

def main(base,workbook,audit,out):
    base,workbook,audit,out=map(Path,[base,workbook,audit,out]);out.mkdir(parents=True,exist_ok=True)
    receipt=json.loads((audit/'receipt.json').read_text())
    assert receipt['status']=='PASS','Do not package a failed audit'
    assert receipt['source_sha256']==sha(workbook),'The audited source changed'
    full=json.loads((audit/'full_source_receipt.json').read_text())
    assert full['status']=='PASS' and full['source_sha256']==sha(workbook),'Full presentation workbook must also pass'
    buyer=out/'buyer';buyer.mkdir(exist_ok=True)
    cache_into_original(workbook,audit/'full_source_recalculated.xlsx',buyer/'Freelancer_Pricing_Intelligence_System_v2_2.xlsx')
    build_guide(base/'Freelancer_Pricing_Intelligence_System_v2_1_Quick_Start.pdf',buyer/'Freelancer_Pricing_Intelligence_System_v2_2_Quick_Start.pdf')
    for name in ['LICENSE.txt','START_HERE.txt']:
        text=(base/name).read_text().replace('v2.1','v2.2').replace('Version: 2.1','Version: 2.2')
        if name=='START_HERE.txt':text+='\nv2.2: enter explicit zeros for no costs/fees/hidden time; incomplete records are excluded from affected averages. Export only CLIENT SUMMARY as PDF. Formula protection is password-free and prevents accidental edits; it is not a privacy or security barrier.\n'
        (buyer/name).write_text(text)
    w=openpyxl.load_workbook(workbook);m=w['MARKET BENCHMARK']
    lines=['DP-FREELANCE-001 v2.2 - CONTEXTUAL MARKET REFERENCES','Checked: 2026-10-06','Historical worldwide contracts, not guaranteed prices or Chile-specific wage recommendations.','']
    for r in range(16,25):lines.extend([f'{m[f"A{r}"].value}: USD {m[f"B{r}"].value}-{m[f"C{r}"].value}/hr',m[f'D{r}'].value,''])
    (buyer/'BENCHMARK_SOURCES.txt').write_text('\n'.join(lines))
    manifest={'product_id':'DP-FREELANCE-001','version':'2.2','supersedes_unreleased':'2.1','state':'AUDIT_PASS_PENDING_UPLOAD_AND_DELIVERY','audit_cases':len(receipt['cases']),'audit_checks_passed':receipt['passed'],'audit_checks_failed':receipt['failed'],'calculation_engine':receipt['engine'],'native_excel_ui_tested':False,'google_sheets_compatibility':'UNVALIDATED','files':{p.name:sha(p) for p in buyer.iterdir() if p.is_file()}}
    (buyer/'DP001_v2_2_MANIFEST.json').write_text(json.dumps(manifest,indent=2))
    archive=out/'Freelancer_Pricing_Intelligence_System_v2_2_BUYER_PACKAGE.zip'
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted(buyer.iterdir()):z.write(p,p.name)
    with zipfile.ZipFile(archive) as z:assert z.testzip() is None
    d=openpyxl.load_workbook(buyer/'Freelancer_Pricing_Intelligence_System_v2_2.xlsx',data_only=True)
    assert d['SYSTEM CHECK']['B19'].value=='PASS' and d['DECISION CENTER']['B16'].value=='GO'
    assert not [(s.title,c.coordinate) for s in d for row in s for c in row if c.data_type=='e']
    print(json.dumps({'buyer_package':str(archive),'sha256':sha(archive),'workbook_sha256':sha(buyer/'Freelancer_Pricing_Intelligence_System_v2_2.xlsx'),'guide_pages':len(PdfReader(buyer/'Freelancer_Pricing_Intelligence_System_v2_2_Quick_Start.pdf').pages),'manifest':manifest}))

if __name__=='__main__':main(*sys.argv[1:])
